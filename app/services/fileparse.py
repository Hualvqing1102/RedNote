"""本地文档解析：把拖拽导入的 PDF / DOCX / TXT / Markdown 转成可提炼的文本。

- 文字型 PDF：直接用 PyMuPDF 抽取文字层；
- 扫描型 PDF（抽不到文字）：逐页渲染成图片后用 RapidOCR(中英文)识别；
- DOCX：python-docx 读段落与表格；
- TXT/MD：按 UTF-8(兼容 GB18030)解码。

所有引擎为纯 pip 安装，便于 PyInstaller 一起打进 exe。
"""
from __future__ import annotations

import threading
from pathlib import Path
from typing import Any

SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".md"}
TEXT_PAGE_MIN = 20  # 一整个 PDF 提取到的字符少于该数视为疑似扫描件
MAX_UPLOAD_BYTES = 80 * 1024 * 1024


class FileParseError(Exception):
    """解析失败(消息可直接展示给用户)。"""


# ---------------------------------------------------------------- 入口


def parse_bytes(filename: str | None, data: bytes) -> dict[str, Any]:
    """按扩展名解析上传内容，返回与网页采集一致的 {title, content, source_url}。"""
    name = (filename or "未命名").strip() or "未命名"
    ext = Path(name).suffix.lower()
    if ext == ".txt" or ext == ".md":
        content = _decode_text(data)
    elif ext == ".docx":
        content = _parse_docx(data)
    elif ext == ".pdf":
        content = _parse_pdf(data)
    else:
        raise FileParseError(f"暂不支持「{ext or '未知'}」格式：请使用 PDF / DOCX / TXT / Markdown")
    content = _clean(content)
    if not content:
        raise FileParseError(f"「{name}」未提取到文字（可能是空白文件或无法识别的扫描件）")
    return {
        "title": Path(name).stem[:200],
        "content": content,
        "source_url": "",
        "filename": name,
    }


def check_size(size: int) -> None:
    if size > MAX_UPLOAD_BYTES:
        raise FileParseError(f"文件过大：上限 {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")


# ---------------------------------------------------------------- 文本


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise FileParseError("无法识别文本编码(支持 UTF-8 / GB18030)")


def _clean(text: str) -> str:
    lines = [ln.rstrip() for ln in (text or "").splitlines()]
    out: list[str] = []
    blank = 0
    for ln in lines:
        if not ln.strip():
            blank += 1
            if blank > 1:
                continue
        else:
            blank = 0
        out.append(ln)
    return "\n".join(out).strip()


# ---------------------------------------------------------------- DOCX


def _parse_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise FileParseError("缺少 DOCX 解析组件，无法读取该文件") from exc
    import io

    parts: list[str] = []
    try:
        document = docx.Document(io.BytesIO(data))
        for p in document.paragraphs:
            if p.text and p.text.strip():
                parts.append(p.text.strip())
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)
    except Exception as exc:  # noqa: BLE001
        raise FileParseError("DOCX 解析失败：文件可能已损坏") from exc
    return "\n\n".join(parts)


# ---------------------------------------------------------------- PDF


def _parse_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise FileParseError("缺少 PDF 解析组件，无法读取该文件") from exc
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise FileParseError("PDF 解析失败：文件可能已损坏或加密") from exc
    try:
        if doc.page_count == 0:
            raise FileParseError("PDF 没有任何页面")
        text = "\n\n".join((page.get_text("text") or "").strip() for page in doc)
        if len(text.strip()) < TEXT_PAGE_MIN:
            text = _ocr_pdf(doc)
        return text
    finally:
        doc.close()


_OCR_LOCK = threading.Lock()
_OCR_ENGINE: Any = None


def _ocr_pdf(doc: Any) -> str:
    """逐页渲染成图片交给 OCR(中英文)。识别不出时返回空串。"""
    global _OCR_ENGINE
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:  # pragma: no cover
        raise FileParseError("未安装 OCR 组件，无法识别扫描版 PDF") from exc
    with _OCR_LOCK:
        if _OCR_ENGINE is None:
            _OCR_ENGINE = RapidOCR()
    engine = _OCR_ENGINE
    out: list[str] = []
    import numpy as np

    for page in doc:
        try:
            import fitz

            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))  # 约 144 dpi
            arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            if pix.n == 4:  # 去掉透明通道，只留 RGB
                arr = arr[:, :, :3]
            result, _ = engine(arr)
            if result:
                for item in result:
                    text = str(item[1]).strip()
                    if text:
                        out.append(text)
        except Exception:  # noqa: BLE001 - 单页失败不阻断整份文档
            continue
    return "\n".join(out)
