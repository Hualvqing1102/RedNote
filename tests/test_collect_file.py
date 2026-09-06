"""拖拽导入本地文档的接口测试：TXT/MD/DOCX/PDF(文字层) 与 OCR 兜底。"""
from __future__ import annotations

import io

import pytest

from app.services import fileparse


@pytest.fixture(autouse=True)
def _ensure_parse_libs():
    pytest.importorskip("docx")
    pytest.importorskip("fitz")


def test_upload_txt(client):
    resp = client.post(
        "/api/collect/file",
        files={"file": ("论文摘要.txt", "这是一段中文摘要。\n\n第二段讲注意力机制。".encode("utf-8"), "text/plain")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["title"] == "论文摘要"
    assert "注意力机制" in body["content"]
    assert body["source_url"] == ""
    assert body["filename"] == "论文摘要.txt"


def test_upload_markdown(client):
    resp = client.post(
        "/api/collect/file",
        files={"file": ("notes.md", b"# Title\n\nHello **bold** text.", "text/markdown")},
    )
    assert resp.status_code == 200
    assert resp.json()["title"] == "notes"
    assert "Hello" in resp.json()["content"]


def test_upload_docx(client):
    import docx

    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("第一段：文档解析测试")
    doc.add_paragraph("第二段：用于验证 DOCX 上传")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    doc.save(buf)
    resp = client.post(
        "/api/collect/file",
        files={"file": ("paper.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["title"] == "paper"
    assert "文档解析测试" in body["content"]
    assert "A | B" in body["content"]


def test_upload_pdf_text_layer(client):
    import fitz

    buf = io.BytesIO()
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Attention Is All You Need and self-attention")
    doc.save(buf)
    resp = client.post("/api/collect/file", files={"file": ("transformer.pdf", buf.getvalue(), "application/pdf")})
    assert resp.status_code == 200
    body = resp.json()
    assert body["title"] == "transformer"
    assert "self-attention" in body["content"]


def test_upload_scanned_pdf_uses_ocr(client, monkeypatch):
    import fitz

    called = {}

    def fake_ocr(doc):
        called["pages"] = doc.page_count
        return "OCR RESULT: scanned page text"

    monkeypatch.setattr(fileparse, "_ocr_pdf", fake_ocr)
    buf = io.BytesIO()
    doc = fitz.open()
    doc.new_page()  # 空白页：无文字层 → 应触发 OCR
    doc.save(buf)
    resp = client.post("/api/collect/file", files={"file": ("scan.pdf", buf.getvalue(), "application/pdf")})
    assert resp.status_code == 200
    assert called.get("pages") == 1
    assert "OCR RESULT" in resp.json()["content"]


def test_upload_unsupported_extension(client):
    resp = client.post("/api/collect/file", files={"file": ("paper.doc", b"old format", "application/msword")})
    assert resp.status_code == 400
    assert "不支持" in resp.json()["detail"]


def test_upload_empty_text_rejected(client):
    resp = client.post("/api/collect/file", files={"file": ("empty.txt", b"   \n  ", "text/plain")})
    assert resp.status_code == 400
    assert "未提取到文字" in resp.json()["detail"]
